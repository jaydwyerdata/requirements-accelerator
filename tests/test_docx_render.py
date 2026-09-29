"""Tests for engine/docx_render.py (e13): the Word (.docx) counterpart to engine/render.py's
plain-text output. These check the same content decisions render.py's own tests already verify
(the layer-00 gate, the readiness gap attribution, the definitions-per-term loop, the reporting
appendix's primary_measure gate), plus the two things unique to this module: the letterhead and
the provenance badge colours. Reads documents back through python-docx's own object model
(paragraphs, tables, cell shading), never the raw XML, except for the one recipe
(_shade_cell/_badge) this module has to hand-roll because python-docx exposes no cell-fill API.
"""

import unittest

from docx.oxml.ns import qn

from engine.docx_render import (
    _NOT_READY_COLOUR,
    _READY_COLOUR,
    _STATE_COLOURS,
    render_business_ask_docx,
    render_technical_spec_docx,
)
from engine.ledger import Ledger, ProvenanceState
from engine.render import LAYER_FIELDS


def _paragraph_texts(doc) -> list[str]:
    # doc.paragraphs only walks the document body, never into table cells, so this never
    # collides with a field value that happens to match a heading's wording.
    return [p.text for p in doc.paragraphs]


def _find_row(doc, field_id: str):
    """The row, across every table in the document, whose first cell is exactly this field id.
    Used to locate a specific field's line in one of the per-layer tables without assuming which
    table index it lands at.
    """
    for table in doc.tables:
        for row in table.rows:
            if row.cells[0].text == field_id:
                return row
    return None


def _cell_fill(cell) -> str | None:
    """The cell's shaded background colour, read back from the w:shd element _shade_cell wrote.
    Not exposed as a public python-docx property, hence the direct OXML lookup."""
    tc_pr = cell._tc.find(qn("w:tcPr"))
    if tc_pr is None:
        return None
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        return None
    return shd.get(qn("w:fill"))


def _fully_stated_ledger() -> Ledger:
    """Every field in every layer, STATED, so nothing is a blocking gap: the readiness badge
    this produces should read READY TO SIZE."""
    ledger = Ledger()
    for layer_id, fields in LAYER_FIELDS.items():
        for field_id in fields:
            ledger.set(field_id, layer_id, f"value for {field_id}", ProvenanceState.STATED,
                       source="test fixture")
    return ledger


class LetterheadTest(unittest.TestCase):
    def test_business_ask_letterhead_names_document_type_and_interview_id(self):
        doc = render_business_ask_docx(Ledger(), interview_id="abcdef12-3456-7890-abcd-ef1234567890")
        texts = _paragraph_texts(doc)
        self.assertIn("Requirements Accelerator", texts)
        self.assertTrue(any("Business Ask" in t and "abcdef12" in t for t in texts))

    def test_technical_spec_letterhead_names_document_type_and_interview_id(self):
        doc = render_technical_spec_docx(Ledger(), interview_id="abcdef12-3456-7890-abcd-ef1234567890")
        texts = _paragraph_texts(doc)
        self.assertIn("Requirements Accelerator", texts)
        self.assertTrue(any("Technical Spec" in t and "abcdef12" in t for t in texts))


class BusinessAskDocxTest(unittest.TestCase):
    def test_no_content_beats_when_problem_not_established(self):
        doc = render_business_ask_docx(Ledger())
        texts = _paragraph_texts(doc)
        self.assertTrue(any("hasn't yet established the problem" in t for t in texts))
        self.assertNotIn("What we understood", texts)

    def test_renders_three_beats_once_problem_is_established(self):
        ledger = Ledger()
        ledger.set("problem_statement", "00", "sales can't see weekly returns",
                   ProvenanceState.STATED, source="requester")
        doc = render_business_ask_docx(ledger)
        texts = _paragraph_texts(doc)
        self.assertIn("What we understood", texts)
        self.assertIn("What you're trying to solve", texts)
        self.assertIn("What happens next", texts)

    def test_ready_message_when_no_requester_gaps_remain(self):
        ledger = Ledger()
        for field_id in ("problem_statement", "query_type", "primary_measure", "grain"):
            ledger.set(field_id, "00", "established", ProvenanceState.STATED, source="requester")
        doc = render_business_ask_docx(ledger)
        self.assertIn("This is ready to size and build.", _paragraph_texts(doc))

    def test_requester_gaps_render_as_bullets(self):
        ledger = Ledger()
        ledger.set("problem_statement", "00", "sales can't see weekly returns",
                   ProvenanceState.STATED, source="requester")
        doc = render_business_ask_docx(ledger)
        bullet_texts = [p.text for p in doc.paragraphs if p.style.name == "List Bullet"]
        self.assertTrue(any("kind of answer" in t for t in bullet_texts))
        self.assertTrue(any("one number" in t for t in bullet_texts))
        self.assertTrue(any("one row" in t for t in bullet_texts))


class TechnicalSpecReadinessTest(unittest.TestCase):
    def test_not_ready_badge_on_an_empty_ledger(self):
        doc = render_technical_spec_docx(Ledger())
        status_cell = doc.tables[0].rows[0].cells[0]
        self.assertEqual(status_cell.text, "NOT READY TO SIZE")
        self.assertEqual(_cell_fill(status_cell), _NOT_READY_COLOUR)

    def test_ready_badge_once_every_field_is_stated(self):
        doc = render_technical_spec_docx(_fully_stated_ledger())
        status_cell = doc.tables[0].rows[0].cells[0]
        self.assertEqual(status_cell.text, "READY TO SIZE")
        self.assertEqual(_cell_fill(status_cell), _READY_COLOUR)


class ProvenanceBadgeTest(unittest.TestCase):
    def test_missing_field_gets_the_missing_badge(self):
        doc = render_technical_spec_docx(Ledger())
        row = _find_row(doc, "problem_statement")
        self.assertIsNotNone(row)
        self.assertEqual(row.cells[2].text, "MISSING")
        self.assertEqual(_cell_fill(row.cells[2]), _STATE_COLOURS[ProvenanceState.MISSING])

    def test_stated_field_gets_the_stated_badge(self):
        ledger = Ledger()
        ledger.set("problem_statement", "00", "sales can't see weekly returns",
                   ProvenanceState.STATED, source="requester")
        doc = render_technical_spec_docx(ledger)
        row = _find_row(doc, "problem_statement")
        self.assertEqual(row.cells[2].text, "STATED")
        self.assertEqual(_cell_fill(row.cells[2]), _STATE_COLOURS[ProvenanceState.STATED])

    def test_inferred_and_assumed_badges_use_their_own_colours(self):
        ledger = Ledger()
        ledger.set("frequency", "00", "weekly", ProvenanceState.INFERRED, source="reasoner")
        ledger.set("blast_radius", "00", "one team", ProvenanceState.ASSUMED, source="reasoner")
        doc = render_technical_spec_docx(ledger)
        inferred_row = _find_row(doc, "frequency")
        assumed_row = _find_row(doc, "blast_radius")
        self.assertEqual(_cell_fill(inferred_row.cells[2]), _STATE_COLOURS[ProvenanceState.INFERRED])
        self.assertEqual(_cell_fill(assumed_row.cells[2]), _STATE_COLOURS[ProvenanceState.ASSUMED])


class DefinitionsSectionTest(unittest.TestCase):
    def test_no_terms_surfaced(self):
        doc = render_technical_spec_docx(Ledger())
        self.assertIn("No ambiguous terms surfaced.", _paragraph_texts(doc))

    def test_one_term_renders_its_definition_and_lock_status(self):
        ledger = Ledger()
        ledger.set("metric_definition::active customer", "07", "purchased in last 90 days",
                   ProvenanceState.STATED, source="requester")
        ledger.set("term_locked::active customer", "07", False, ProvenanceState.STATED,
                   source="requester", note="exclusions not yet confirmed")
        doc = render_technical_spec_docx(ledger)
        texts = _paragraph_texts(doc)
        self.assertTrue(any('"active customer"' in t and "not locked" in t for t in texts))
        self.assertTrue(any("purchased in last 90 days" in t for t in texts))
        self.assertTrue(any("exclusions not yet confirmed" in t for t in texts))


class ReportingSpecTest(unittest.TestCase):
    def test_omitted_without_a_primary_measure(self):
        doc = render_technical_spec_docx(Ledger())
        self.assertNotIn("Reporting specification", _paragraph_texts(doc))

    def test_included_once_a_primary_measure_exists(self):
        ledger = Ledger()
        ledger.set("primary_measure", "02", "weekly return count", ProvenanceState.STATED,
                   source="requester")
        doc = render_technical_spec_docx(ledger)
        self.assertIn("Reporting specification", _paragraph_texts(doc))
        measure_row = next(
            row for table in doc.tables for row in table.rows if row.cells[0].text == "Measure"
        )
        self.assertEqual(measure_row.cells[1].text, "weekly return count")


if __name__ == "__main__":
    unittest.main()
