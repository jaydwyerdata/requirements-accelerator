"""e13: renders the two output documents as Word (.docx) documents, the presentable counterpart
to engine/render.py's plain-text strings. docs/output-template.md already defines what content
goes where; this module is purely the presentation layer on top of it, so it imports the same
field lists and gap logic from engine/render.py rather than re-deriving them, the one thing that
must never drift out of sync between the two renderers.

Why a separate module rather than teaching render.py to emit docx: render.py's plain-text output
still has a live job (cli.py's terminal printout, app.py's on-screen display once an interview
finishes), so it stays exactly as it is. This module is delivery-only output, used by
engine/delivery.py, and it is the one place in the engine that imports python-docx, the same
import-isolation reasoning engine/cortex_reasoner.py gives for keeping its own dependency (there,
snowflake.snowpark) out of the shared modules.

Design agreed 14 September 2026 (see docs/architecture-decisions.md's e13 entry): a light
letterhead on both documents (title, document type, interview id, date), and provenance state
shown as a coloured badge on the technical spec, since that is the single clearest demo moment
this tool has, a plain-language answer sitting next to its engineer-grade translation with the
gap visually obvious rather than something you have to read carefully to notice.
"""

from datetime import datetime, timezone

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .ledger import Ledger, ProvenanceState, owner_for
from .render import (
    ALL_FIELDS,
    BLOCKING_FIELDS,
    GAP_PLAIN_LANGUAGE,
    GAP_REASONS,
    LAYER_FIELDS,
    LAYER_NAMES,
    _terms_in_ledger,
)

# Same palette the build tracker artifact uses (--done, --doing, --prio), plus one added blue for
# INFERRED, so a badge colour means the same thing everywhere this project shows status.
_STATE_COLOURS = {
    ProvenanceState.STATED: "1F7A4D",
    ProvenanceState.INFERRED: "2B6CB0",
    ProvenanceState.ASSUMED: "9C6B0C",
    ProvenanceState.MISSING: "B23A2C",
}
_READY_COLOUR = "1F7A4D"
_NOT_READY_COLOUR = "B23A2C"
_MUTED_GREY = "57615E"


def _shade_cell(cell, hex_colour: str) -> None:
    """Sets a table cell's background fill. Not exposed as a public python-docx API, this is the
    standard OXML recipe for it (add a w:shd element to the cell's tcPr)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_colour)
    tc_pr.append(shd)


def _badge(cell, text: str, hex_colour: str) -> None:
    """Turns a table cell into a coloured badge: shaded background, bold white centred text.
    Used for the provenance state column in the technical spec, the one visual element this
    template was specifically built around."""
    _shade_cell(cell, hex_colour)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def _letterhead(doc: Document, document_type: str, ledger: Ledger, interview_id: str) -> None:
    doc.add_heading("Requirements Accelerator", level=0)
    subtitle = doc.add_paragraph()
    short_id = interview_id[:8] if interview_id else "unknown"
    # Built from parts rather than strftime("%-d ..."): the "-" no-padding flag is a glibc
    # extension, and Windows' strftime rejects it with ValueError: Invalid format string.
    # Caught on the first full local run on Windows (demo recording, 7 October 2026).
    now = datetime.now(timezone.utc)
    date_str = f"{now.day} {now:%B %Y}"
    run = subtitle.add_run(f"{document_type}  •  Interview {short_id}  •  {date_str}")
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(_MUTED_GREY)
    doc.add_paragraph()


def render_business_ask_docx(ledger: Ledger, interview_id: str = "") -> Document:
    """Word counterpart to engine.render.render_business_ask: same three beats, same gate on
    layer 00 never having cleared, but as a letterheaded document a requester would actually
    keep, not a wall of plain text."""
    doc = Document()
    _letterhead(doc, "Business Ask", ledger, interview_id)

    problem = ledger.get("problem_statement")
    if problem is None or problem.state == ProvenanceState.MISSING:
        doc.add_paragraph(
            "This interview hasn't yet established the problem, so there is nothing genuine "
            "to reflect back yet."
        )
        return doc

    process = ledger.get("current_process")
    activation = ledger.get("activation_use_case")
    acceptance = ledger.get("acceptance_criteria")

    doc.add_heading("What we understood", level=1)
    doc.add_paragraph(str(problem.value))
    if process is not None:
        doc.add_paragraph(f"Today, this is handled by: {process.value}")

    doc.add_heading("What you're trying to solve", level=1)
    if activation is not None:
        doc.add_paragraph(str(activation.value))
    if acceptance is not None:
        doc.add_paragraph(f"You'll know it's fixed when: {acceptance.value}")

    doc.add_heading("What happens next", level=1)
    requester_gaps = []
    for field_id in ("query_type", "primary_measure", "grain"):
        record = ledger.get(field_id)
        if record is None or record.state == ProvenanceState.MISSING:
            requester_gaps.append(GAP_PLAIN_LANGUAGE[field_id])
    for term in _terms_in_ledger(ledger):
        locked = ledger.get(f"term_locked::{term}")
        if locked is None or not locked.value:
            requester_gaps.append(f'we\'d still like to pin down exactly what "{term}" means')

    if requester_gaps:
        for gap in requester_gaps:
            doc.add_paragraph(gap, style="List Bullet")
    else:
        doc.add_paragraph("Nothing more is needed from you for now. The data team will review this and come back to you.")

    return doc


def _add_field_table(doc: Document, layer_id: str, fields: list[str], ledger: Ledger) -> None:
    doc.add_heading(LAYER_NAMES[layer_id], level=2)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    header = table.rows[0].cells
    for cell, text in zip(header, ("Field", "Value", "State", "Source")):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True

    for field_id in fields:
        record = ledger.get(field_id)
        row = table.add_row().cells
        row[0].text = field_id
        if record is None or record.state == ProvenanceState.MISSING:
            row[1].text = "--"
            _badge(row[2], "MISSING", _STATE_COLOURS[ProvenanceState.MISSING])
            row[3].text = f"owner: {owner_for(field_id).value}"
        else:
            value_text = str(record.value)
            if record.note:
                value_text += f"  ({record.note})"
            row[1].text = value_text
            _badge(row[2], record.state.name, _STATE_COLOURS[record.state])
            row[3].text = record.source


def _add_readiness_block(doc: Document, ledger: Ledger) -> None:
    non_missing = 0
    blocking_gaps: list[tuple[str, str, str]] = []
    workable_gaps: list[str] = []

    for field_id in ALL_FIELDS:
        record = ledger.get(field_id)
        if record is not None and record.state != ProvenanceState.MISSING:
            non_missing += 1
            continue
        if field_id in BLOCKING_FIELDS:
            blocking_gaps.append(
                (field_id, GAP_REASONS.get(field_id, "not yet established"),
                 owner_for(field_id).value)
            )
        else:
            workable_gaps.append(field_id)

    total_fields = len(ALL_FIELDS)
    for term in _terms_in_ledger(ledger):
        total_fields += 1
        locked = ledger.get(f"term_locked::{term}")
        if locked is not None and locked.value:
            non_missing += 1
        else:
            reason = locked.note if (locked is not None and locked.note) else "not yet locked"
            blocking_gaps.append((f"term: {term}", reason, "requester"))

    completion = round(100 * non_missing / total_fields) if total_fields else 0
    ready = not blocking_gaps

    doc.add_heading("Readiness", level=1)
    status_table = doc.add_table(rows=1, cols=1)
    status_cell = status_table.rows[0].cells[0]
    _badge(
        status_cell,
        "READY TO SIZE" if ready else "NOT READY TO SIZE",
        _READY_COLOUR if ready else _NOT_READY_COLOUR,
    )
    doc.add_paragraph(f"Completion: {completion}% ({non_missing}/{total_fields} fields)")

    if blocking_gaps:
        doc.add_paragraph(f"Blocking gaps ({len(blocking_gaps)}):")
        gap_table = doc.add_table(rows=1, cols=3)
        gap_table.style = "Light Grid Accent 1"
        header = gap_table.rows[0].cells
        for cell, text in zip(header, ("Field", "Reason", "Owner")):
            cell.text = text
            cell.paragraphs[0].runs[0].bold = True
        for field_id, reason, owner in blocking_gaps:
            row = gap_table.add_row().cells
            row[0].text = field_id
            row[1].text = reason
            row[2].text = owner

    doc.add_paragraph(f"Workable gaps not yet collected: {len(workable_gaps)}")
    doc.add_paragraph()


def _add_definitions_section(doc: Document, ledger: Ledger) -> None:
    doc.add_heading("Definitions", level=1)
    terms = _terms_in_ledger(ledger)
    if not terms:
        doc.add_paragraph("No ambiguous terms surfaced.")
        return

    for term in terms:
        definition = ledger.get(f"metric_definition::{term}")
        competing = ledger.get(f"competing_definition_check::{term}")
        exclusions = ledger.get(f"exclusion_filters::{term}")
        synonyms = ledger.get(f"synonyms::{term}")
        locked = ledger.get(f"term_locked::{term}")
        status = "locked" if (locked is not None and locked.value) else "not locked"

        doc.add_heading(f'"{term}" -- {status}', level=2)
        if definition is not None:
            doc.add_paragraph(f"Definition: {definition.value}")
        if exclusions is not None:
            doc.add_paragraph(f"Excludes: {exclusions.value}")
        if synonyms is not None:
            doc.add_paragraph(f"Also called: {synonyms.value}")
        if competing is not None:
            doc.add_paragraph(f"Competing definition check: {competing.value}")
        if locked is not None and not locked.value and locked.note:
            doc.add_paragraph(f"Gap: {locked.note}")


def _add_reporting_spec(doc: Document, ledger: Ledger) -> None:
    primary_measure = ledger.get("primary_measure")
    if primary_measure is None or primary_measure.state == ProvenanceState.MISSING:
        return

    dimensions = ledger.get("dimensions")
    drill = ledger.get("drill_through_required")
    grain = ledger.get("grain")
    benchmark = ledger.get("benchmark_comparison")
    freshness = ledger.get("freshness_sla")
    schedule = ledger.get("schedule_alignment")
    rls = ledger.get("row_level_security")

    terms = _terms_in_ledger(ledger)
    filter_parts = []
    for term in terms:
        exclusions = ledger.get(f"exclusion_filters::{term}")
        if exclusions is not None:
            filter_parts.append(f"{term}: {exclusions.value}")
    filters_text = "; ".join(filter_parts) if filter_parts else "not established"

    refresh_text = freshness.value if freshness is not None else "not established"
    if schedule is not None and str(schedule.value).strip().lower() not in ("", "none"):
        refresh_text = f"{refresh_text}, aligned to {schedule.value}"

    doc.add_heading("Reporting specification", level=1)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    rows = [
        ("Measure", str(primary_measure.value)),
        ("Dimensions", dimensions.value if dimensions is not None else "not established"),
        ("Drill path",
         f"{drill.value if drill is not None else 'not established'} "
         f"(grain: {grain.value if grain is not None else 'not established'})"),
        ("Comparison basis",
         benchmark.value if benchmark is not None else "not established"),
        ("Filters", filters_text),
        ("Refresh cadence", refresh_text),
        ("Row-level security", rls.value if rls is not None else "not established"),
    ]
    for label, value in rows:
        row = table.add_row().cells
        row[0].text = label
        row[0].paragraphs[0].runs[0].bold = True
        row[1].text = str(value)


def render_technical_spec_docx(ledger: Ledger, interview_id: str = "") -> Document:
    """Word counterpart to engine.render.render_technical_spec. Readiness leads, same as the
    plain-text version, then layers 00-09, definitions, and the reporting appendix. The one
    real addition over the plain-text renderer is the provenance badge on every field row, see
    this module's docstring for why."""
    doc = Document()
    _letterhead(doc, "Technical Spec (internal only)", ledger, interview_id)

    _add_readiness_block(doc, ledger)
    for layer_id, fields in LAYER_FIELDS.items():
        _add_field_table(doc, layer_id, fields, ledger)
    _add_definitions_section(doc, ledger)
    _add_reporting_spec(doc, ledger)

    return doc
